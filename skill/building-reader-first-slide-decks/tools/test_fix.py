import os
import tempfile
import unittest
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt

from fix import add_links, check_cites, check_pptx, merge_cites, check_source, embed_fonts, eot, fix_pptx, normalize_radicals, pdf_font_problems

FONTS = [f"node_modules/@expo-google-fonts/noto-sans-jp/{d}/NotoSansJP_{d}.ttf" for d in ("400Regular", "700Bold")]


def read_slide_xml(path, n):
    return zipfile.ZipFile(path).read(f"ppt/slides/slide{n}.xml").decode("utf8")


def make_raw(path):
    """Marpの変換直後を模したファイル: 部首用の文字、特殊なフォント名、10pt、中国語、句点"""
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[6])
    for text in ["⽅式A", "終わり。"]:
        r = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(10)
        r.font.name = "NotoSansCJKjp"
        r._r.get_or_add_rPr().set("lang", "zh-CN")
    p.save(path)


class FixTest(unittest.TestCase):
    def setUp(self):
        d = tempfile.mkdtemp()
        self.raw, self.out = os.path.join(d, "raw.pptx"), os.path.join(d, "out.pptx")
        make_raw(self.raw)

    def test_radicals_are_normalized(self):
        self.assertEqual(normalize_radicals("⽅式：⾮root"), "方式：非root")

    def test_check_reports_each_violation(self):
        v = "\n".join(check_pptx(self.raw))
        for word in ["12pt", "句点", "部首", "フォント", "埋め込まれていない"]:
            self.assertIn(word, v)

    def test_fix_clears_radicals_and_font(self):
        fix_pptx(self.raw, self.out)
        v = "\n".join(check_pptx(self.out))
        self.assertNotIn("部首", v)
        self.assertNotIn("以外", v)  # フォント名の違反が消える（埋め込みは別の工程）
        self.assertIn("12pt", v)  # 大きさと句点は原稿側で直すものなので残る
        self.assertIn("句点", v)

    def test_lang_is_japanese(self):
        fix_pptx(self.raw, self.out)
        xml = read_slide_xml(self.out, 1)
        self.assertNotIn('lang="zh-CN"', xml)
        self.assertIn('lang="ja-JP"', xml)

    def test_source_rules(self):
        # タイトルは「# 」、メインメッセージは「## 」。長さの決まりはメインメッセージだけに掛かる
        md = "# " + "題" * 40 + "\n## " + "あ" * 33 + "\n\n---\n\n<!-- _class: cover -->\n## " + "い" * 40 + "\n\n文。\n"
        v = "\n".join(check_source(md))
        self.assertIn("32字", v)
        self.assertEqual(v.count("32字"), 1)  # cover は対象外
        self.assertIn("句点", v)
        # 半角の英数字は半分の幅として数える
        self.assertEqual(check_source("## " + "a" * 12 + "あ" * 26), [])
        self.assertIn("32字", "\n".join(check_source("## " + "a" * 12 + "あ" * 27)))

    def test_pdf_fonts_must_be_noto_sans_jp(self):
        head = "name type encoding emb sub uni object ID\n---- ---- ---\n"
        bad = head + "BAAAAA+WenQuanYiZenHei TrueType WinAnsi yes yes yes 389 0\nCAAAAA+NotoSansJP-Bold CID yes yes yes 1 0\n"
        good = head + "AAAAAA+NotoSansJP-Regular CID yes yes yes 1 0\nCAAAAA+NotoSansJP-Bold CID yes yes yes 2 0\n"
        self.assertEqual(pdf_font_problems(good), [])
        self.assertIn("WenQuanYiZenHei", "\n".join(pdf_font_problems(bad)))
        self.assertIn("埋め込みフォントがない", "\n".join(pdf_font_problems(head)))

    def test_eot_wraps_the_font_unchanged(self):
        import struct
        ttf = open(FONTS[0], "rb").read()
        data = eot(FONTS[0])
        total, size, version, flags = struct.unpack("<IIII", data[:16])
        self.assertEqual((total, size, version, flags), (len(data), len(ttf), 0x00020002, 0))
        self.assertEqual(data[-len(ttf):], ttf)

    def test_embed_adds_fonts_and_tells_powerpoint_to_use_them(self):
        fix_pptx(self.raw, self.out)
        embed_fonts(self.out, *FONTS)
        z = zipfile.ZipFile(self.out)
        self.assertEqual(sorted(n for n in z.namelist() if n.endswith(".fntdata")),
                         ["ppt/fonts/font1.fntdata", "ppt/fonts/font2.fntdata"])
        pres = z.read("ppt/presentation.xml").decode()
        self.assertEqual(pres.count('embedTrueTypeFonts="1"'), 1)  # 属性が重複するとファイルが開けなくなる
        self.assertIn('<p:font typeface="Noto Sans JP"', pres)
        self.assertIn('charset="-128"', pres)
        self.assertIn('Extension="fntdata"', z.read("[Content_Types].xml").decode())
        self.assertEqual(z.read("ppt/_rels/presentation.xml.rels").decode().count("relationships/font"), 2)
        self.assertNotIn("埋め込まれていない", "\n".join(check_pptx(self.out)))
        Presentation(self.out)  # 壊れていなければ開ける

    def test_links_are_restored_across_split_runs(self):
        p = Presentation()
        tf = p.slides.add_slide(p.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame
        for text in ["VS Code ", "公式「", "Sharing", "」", "（補足）"]:
            r = tf.paragraphs[0].add_run()
            r.text = text
            r.font.size = Pt(12)
        p.save(self.out)
        md = '<a href="https://example.com/a?x=1&y=2">VS Code 公式「Sharing」</a>（補足） <a href="https://example.com/none">ない文字</a>'
        self.assertEqual(add_links(self.out, md), ["リンクを戻せなかった: https://example.com/none"])
        runs = Presentation(self.out).slides[0].shapes[0].text_frame.paragraphs[0].runs
        self.assertEqual([r.hyperlink.address for r in runs],
                         ["https://example.com/a?x=1&y=2"] * 4 + [None])

    def test_citation_is_joined_to_its_message_as_superscript(self):
        def sp(x, text, attrs):
            return (f'<p:sp><p:spPr><a:xfrm><a:off x="{x}" y="100"/></a:xfrm></p:spPr><p:txBody><a:p><a:r>'
                    f'<a:rPr {attrs}><a:solidFill><a:srgbClr val="0017c1"/></a:solidFill></a:rPr>'
                    f'<a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>')
        xml = sp(10, "左の根拠", 'b="1" sz="1800"') + sp(5000, "右の根拠", 'b="1" sz="1800"') + sp(9000, "2‒3", 'b="0" sz="1200"')  # 変換後は横線が別の文字になる
        out = merge_cites(xml, {"2–3"})
        self.assertIn("<a:t> 2–3</a:t>", out)
        self.assertEqual(out.count("<p:sp>"), 2)  # 番号の箱はなくなる
        self.assertIn('右の根拠</a:t></a:r><a:r><a:rPr b="0" sz="2100" baseline="30000">', out)
        self.assertNotIn("左の根拠</a:t></a:r><a:r>", out)

    def test_citations_match_references_in_order_of_first_use(self):
        refs = "\n# 参考文献\n<tr><td>1.</td><td>A</td></tr>\n<tr><td>2.</td><td>B</td></tr>\n<tr><td>3.</td><td>C</td></tr>\n"
        self.assertEqual(check_cites("<h3>a<sup>1</sup></h3><h3>b<sup>1-3</sup></h3>" + refs), [])
        v = "\n".join(check_cites("<h3>a<sup>2</sup></h3><h3>b<sup>1, 4</sup></h3>" + refs))
        for word in ["初めて出た順", "引用 4 が参考文献にない", "参考文献 3 がどこからも"]:
            self.assertIn(word, v)


if __name__ == "__main__":
    unittest.main()
